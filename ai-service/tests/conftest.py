import io
import pytest
import docx


def create_sample_pdf(pages: list[str]) -> bytes:
    """Generate valid in-memory multi-page PDF for test fixtures."""
    page_obj_ids = []
    curr_id = 3
    for _ in pages:
        page_obj_ids.append(curr_id)
        curr_id += 2

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}

    def write_obj(oid: int, content: bytes):
        offsets[oid] = out.tell()
        out.write(f"{oid} 0 obj\n".encode("ascii"))
        out.write(content)
        out.write(b"\nendobj\n")

    # Catalog & Pages
    write_obj(1, b"<</Type/Catalog/Pages 2 0 R>>")
    kids_str = " ".join(f"{pid} 0 R" for pid in page_obj_ids)
    write_obj(2, f"<</Type/Pages/Count {len(pages)}/Kids[{kids_str}]>>".encode("ascii"))

    font_id = curr_id
    curr_id += 1

    for idx, text in enumerate(pages):
        pid = page_obj_ids[idx]
        cid = pid + 1

        stream_lines = ["BT", "/F1 12 Tf", "50 750 Td"]
        for line in text.split("\n"):
            clean_l = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream_lines.append(f"({clean_l}) Tj")
            stream_lines.append("0 -18 Td")
        stream_lines.append("ET")
        stream_bytes = "\n".join(stream_lines).encode("latin-1")

        write_obj(
            pid,
            f"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Resources<</Font<</F1 {font_id} 0 R>>>>/Contents {cid} 0 R>>".encode(
                "ascii"
            ),
        )
        write_obj(
            cid,
            f"<</Length {len(stream_bytes)}>>\nstream\n".encode("ascii")
            + stream_bytes
            + b"\nendstream",
        )

    write_obj(font_id, b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>")

    startxref = out.tell()
    out.write(f"xref\n0 {curr_id}\n".encode("ascii"))
    out.write(b"0000000000 65535 f \n")
    for oid in range(1, curr_id):
        out.write(f"{offsets[oid]:010d} 00000 n \n".encode("ascii"))
    out.write(
        f"trailer<</Size {curr_id}/Root 1 0 R>>\nstartxref\n{startxref}\n%%EOF\n".encode("ascii")
    )
    return out.getvalue()


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    pages = [
        "1. MASTER SERVICES AGREEMENT\nThis Agreement is made between Client and Vendor.",
        "2. CONFIDENTIALITY\nBoth parties agree to protect proprietary materials.\n\n3. INDEMNIFICATION\nVendor shall indemnify Client against third-party claims.",
    ]
    return create_sample_pdf(pages)


@pytest.fixture
def sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("1. NON-DISCLOSURE AGREEMENT", level=1)
    doc.add_paragraph("This Non-Disclosure Agreement governs confidential disclosures.")

    doc.add_heading("2. OBLIGATIONS", level=2)
    doc.add_paragraph("The Recipient agrees not to disclose Confidential Information.")
    doc.add_paragraph("• Recipient will protect information with reasonable care.")
    doc.add_paragraph("• Recipient will not reverse engineer any materials.")

    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Term Field"
    table.cell(0, 1).text = "Contract Specification"
    table.cell(1, 0).text = "Effective Date"
    table.cell(1, 1).text = "January 1, 2026"
    table.cell(2, 0).text = "Jurisdiction"
    table.cell(2, 1).text = "Delaware, USA"

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.fixture
def sample_txt_content() -> str:
    return """EMPLOYMENT CONTRACT

1. POSITION AND DUTIES
The Employee agrees to perform duties faithfully and to the best of their abilities.

2. COMPENSATION AND BENEFITS
The Company shall pay Employee an annual base salary payable in bi-weekly increments.
    (a) The Employee is eligible for an annual discretionary bonus.
    (b) The Employee is entitled to 20 business days of paid vacation.
        (i) Unused vacation may not roll over past March 31.

3. TERMINATION
Either party may terminate this employment relationship at will with 30 days written notice.
"""
